# -*- coding: utf-8 -*-
"""
_00_corpus_overview.py
======================

Copyright © 2026 Till Preidt (GitHub: MrNemesis98).
This work is licensed under the Creative Commons Attribution–NonCommercial 4.0 International License
(CC BY-NC 4.0). Reuse and adaptation are permitted for non-commercial purposes,
provided appropriate credit is given.

Purpose
-------
Creates an XLSX overview for one corpus subset and links the source audiobook
folders to the Whisper segment WAV files prepared for MFA.

Input
-----
data/A___preprocessing/00___data_collections/Wolnelektury/<subset>/
data/A___preprocessing/03___whisper_output/input_for_mfa/<subset>/

Output
------
data/B___statistics/00___corpus_overview_output/<subset>/
    corpus_overview_<subset>.xlsx
    audio_duration_cache_<subset>.json

The subset is selected through DATASET_SUFFIX and must be one of:
training, validation, or test.

Requirements
------------
    python -m pip install openpyxl pydub

For MP3 duration detection, pydub needs FFmpeg available on PATH. WAV duration
can be read via the Python standard library.

Performance note
----------------
MP3 duration detection can be slow because FFmpeg has to inspect many files.
The script therefore stores durations in a subset-specific JSON cache. The
first run may still take a while; later runs should be much faster. Use
--skip-durations for a quick overview without duration calculation, or
--refresh-cache to force recalculation.

Author
------
Prepared for the Fricavis / Polish fricative analysis workflow.
"""

from __future__ import annotations

import argparse
import json
import math
import os
import re
import sys
import wave
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Set, Tuple


# ---------------------------------------------------------------------------
# Central project configuration
# ---------------------------------------------------------------------------
SCRIPT_DIR = Path(__file__).resolve().parent
SRC_DIR = SCRIPT_DIR.parent

if str(SRC_DIR) not in sys.path:
    sys.path.insert(0, str(SRC_DIR))

import _00_config as config  # noqa: E402


try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError as exc:
    raise SystemExit(
        "ERROR: openpyxl is not installed.\n"
        "Install it with:\n"
        "    python -m pip install openpyxl\n"
    ) from exc


# ---------------------------------------------------------------------------
# Script-specific settings
# ---------------------------------------------------------------------------
DATASET_SUFFIX = "training"  # allowed: training, validation, test
CORPUS_SOURCE_FOLDER = "Wolnelektury"
AUDIO_EXTENSIONS = {".wav", ".wave", ".mp3"}
WAV_EXTENSIONS = {".wav", ".wave"}
MP3_EXTENSIONS = {".mp3"}
TEXT_EXTENSIONS = {".txt", ".lab"}
SEGMENT_PATTERN = re.compile(r"^(?P<base>.+)_seg_(?P<idx>\d+)\.wav$", re.IGNORECASE)


@dataclass
class AudiobookStats:
    folder_name: str
    title: str
    author: str
    folder_path: Path
    chapter_count: int
    total_duration_sec: float
    audio_format: str
    possible_prefixes: Set[str]
    matched_units: int
    total_segments: int
    segments_per_chapter: Optional[float]
    status: str
    missing_units_preview: str


@dataclass
class DurationOptions:
    calculate_durations: bool = True
    use_cache: bool = True
    refresh_cache: bool = False


def parse_title_author(folder_name: str) -> Tuple[str, str]:
    """Split folder name at the last comma: 'Title, Author'."""
    if "," not in folder_name:
        return folder_name.strip(), ""
    title, author = folder_name.rsplit(",", 1)
    return title.strip(), author.strip()


def iter_files_by_extensions(folder: Path, extensions: Set[str]) -> List[Path]:
    """Return files with matching extensions, recursively, in stable order."""
    files: List[Path] = []
    for path in folder.rglob("*"):
        if path.is_file() and path.suffix.lower() in extensions:
            files.append(path)
    return sorted(files, key=lambda p: str(p.relative_to(folder)).lower())


def make_cache_key(path: Path) -> str:
    """Return a stable cache key including path, size, and modification time."""
    try:
        stat = path.stat()
        return f"{str(path.resolve())}|{stat.st_size}|{int(stat.st_mtime)}"
    except OSError:
        return str(path.resolve())


def load_duration_cache(cache_path: Path) -> Dict[str, float]:
    """Load cached duration values from JSON."""
    if not cache_path.exists():
        return {}
    try:
        data = json.loads(cache_path.read_text(encoding="utf-8"))
        if isinstance(data, dict):
            return {str(k): float(v) for k, v in data.items()}
    except Exception:
        pass
    return {}


def save_duration_cache(cache_path: Path, cache: Dict[str, float]) -> None:
    """Save cached duration values to JSON."""
    cache_path.parent.mkdir(parents=True, exist_ok=True)
    cache_path.write_text(json.dumps(cache, ensure_ascii=False, indent=2, sort_keys=True), encoding="utf-8")


def wav_duration_sec(path: Path) -> float:
    """Return WAV duration in seconds using the standard library."""
    try:
        with wave.open(str(path), "rb") as wf:
            frames = wf.getnframes()
            rate = wf.getframerate()
            if rate <= 0:
                return 0.0
            return frames / float(rate)
    except Exception:
        # Keep the overview robust: one broken WAV should not stop the report.
        return 0.0


def mp3_duration_sec(path: Path) -> float:
    """Return MP3 duration in seconds using pydub/FFmpeg."""
    try:
        from pydub import AudioSegment
        return len(AudioSegment.from_file(path)) / 1000.0
    except Exception:
        # Keep the overview robust: one broken MP3 or missing FFmpeg should not stop the report.
        return 0.0


def audio_duration_sec(path: Path, cache: Dict[str, float], options: DurationOptions) -> float:
    """Return audio duration for supported source formats, optionally using cache."""
    if not options.calculate_durations:
        return 0.0

    key = make_cache_key(path)
    if options.use_cache and not options.refresh_cache and key in cache:
        return cache[key]

    suffix = path.suffix.lower()
    if suffix in WAV_EXTENSIONS:
        duration = wav_duration_sec(path)
    elif suffix in MP3_EXTENSIONS:
        duration = mp3_duration_sec(path)
    else:
        duration = 0.0

    if options.use_cache:
        cache[key] = float(duration)
    return duration


def deduplicate_audio_files(audio_files: List[Path]) -> List[Path]:
    """Avoid double-counting chapters when MP3 and WAV versions share a stem.

    If the same chapter exists as both WAV and MP3, prefer the WAV file because
    it is usually the preprocessed version. Otherwise, keep all unique stems.
    """
    by_stem: Dict[str, List[Path]] = defaultdict(list)
    for path in audio_files:
        by_stem[path.stem.lower()].append(path)

    chosen: List[Path] = []
    for _stem, files in sorted(by_stem.items(), key=lambda item: item[0]):
        wavs = [p for p in files if p.suffix.lower() in WAV_EXTENSIONS]
        if wavs:
            chosen.append(sorted(wavs, key=lambda p: p.name.lower())[0])
        else:
            chosen.append(sorted(files, key=lambda p: p.name.lower())[0])
    return sorted(chosen, key=lambda p: str(p).lower())


def describe_audio_format(audio_files: List[Path]) -> str:
    """Return a short human-readable audio format description."""
    formats = sorted({p.suffix.lower().lstrip(".").upper().replace("WAVE", "WAV") for p in audio_files})
    if not formats:
        return "N/V"
    if len(formats) == 1:
        return formats[0]
    return "mixed (" + ", ".join(formats) + ")"


def format_hhmmss(seconds: float) -> str:
    """Format duration seconds as HH:MM:SS."""
    seconds_int = int(round(seconds))
    hours = seconds_int // 3600
    minutes = (seconds_int % 3600) // 60
    secs = seconds_int % 60
    return f"{hours:02d}:{minutes:02d}:{secs:02d}"


def normalize_prefix(stem: str) -> str:
    """Normalize stems only lightly; segment matching remains exact by default."""
    return stem.strip()


def collect_segment_index(mfa_input_folder: Path) -> Dict[str, Set[str]]:
    """Map segment base prefix -> set of segment WAV filenames."""
    index: Dict[str, Set[str]] = defaultdict(set)
    if not mfa_input_folder.exists():
        return index

    for wav_path in mfa_input_folder.rglob("*.wav"):
        match = SEGMENT_PATTERN.match(wav_path.name)
        if not match:
            continue
        base = normalize_prefix(match.group("base"))
        index[base].add(wav_path.name)

    # Also catch uppercase .WAV if glob above misses platform-specific cases.
    for wav_path in mfa_input_folder.rglob("*.WAV"):
        match = SEGMENT_PATTERN.match(wav_path.name)
        if not match:
            continue
        base = normalize_prefix(match.group("base"))
        index[base].add(wav_path.name)

    return index


def analyze_audiobook_folder(folder: Path, segment_index: Dict[str, Set[str]], duration_cache: Dict[str, float], duration_options: DurationOptions) -> AudiobookStats:
    """Compute one row of the corpus overview."""
    title, author = parse_title_author(folder.name)

    raw_audio_files = iter_files_by_extensions(folder, AUDIO_EXTENSIONS)
    audio_files = deduplicate_audio_files(raw_audio_files)
    text_files = iter_files_by_extensions(folder, TEXT_EXTENSIONS)

    # Chapters are counted as unique source audio stems. MP3 and WAV are both
    # supported; duplicate MP3/WAV versions with identical stems are counted once.
    chapter_count = len(audio_files)
    total_duration_sec = sum(audio_duration_sec(p, duration_cache, duration_options) for p in audio_files)
    audio_format = describe_audio_format(audio_files)

    # Segment names are created from transcript stems in the Whisper segmenter.
    # We therefore include both text and source-audio stems as possible matching prefixes.
    possible_prefixes = {normalize_prefix(p.stem) for p in audio_files + text_files}

    matched_prefixes = sorted(prefix for prefix in possible_prefixes if prefix in segment_index)
    total_segments = sum(len(segment_index[prefix]) for prefix in matched_prefixes)
    matched_units = len(matched_prefixes)

    expected_units = len({normalize_prefix(p.stem) for p in text_files}) if text_files else chapter_count
    if expected_units == 0:
        status = "no source files"
    elif total_segments == 0:
        status = "not segmented yet"
    elif matched_units >= expected_units:
        status = "segmented"
    else:
        status = "partially segmented"

    segments_per_chapter = total_segments / chapter_count if chapter_count else None

    expected_prefixes = {normalize_prefix(p.stem) for p in text_files} if text_files else {normalize_prefix(p.stem) for p in audio_files}
    missing = sorted(expected_prefixes.difference(matched_prefixes))
    missing_preview = ", ".join(missing[:5])
    if len(missing) > 5:
        missing_preview += f" ... (+{len(missing) - 5} more)"

    return AudiobookStats(
        folder_name=folder.name,
        title=title,
        author=author,
        folder_path=folder,
        chapter_count=chapter_count,
        total_duration_sec=total_duration_sec,
        audio_format=audio_format,
        possible_prefixes=possible_prefixes,
        matched_units=matched_units,
        total_segments=total_segments,
        segments_per_chapter=segments_per_chapter,
        status=status,
        missing_units_preview=missing_preview,
    )


def autosize_columns(ws, min_width: int = 8, max_width: int = 48) -> None:
    """Apply reasonable column widths."""
    for column_cells in ws.columns:
        max_len = 0
        col_letter = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            value = cell.value
            if value is None:
                continue
            max_len = max(max_len, len(str(value)))
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, min_width), max_width)


def create_workbook(stats: List[AudiobookStats], collections_folder: Path, mfa_input_folder: Path, output_path: Path, duration_options: DurationOptions, cache_path: Path) -> None:
    """Create formatted XLSX workbook."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Corpus Overview"

    generated_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    # Title block
    ws["A1"] = "Audiobook Corpus Overview"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = f"Generated: {generated_at}"
    ws["A3"] = f"Source folder: {collections_folder}"
    ws["A4"] = f"MFA segment folder: {mfa_input_folder}"
    ws["A5"] = f"Duration calculation: {'enabled' if duration_options.calculate_durations else 'skipped'} | Cache: {'enabled' if duration_options.use_cache else 'disabled'} | Cache file: {cache_path}"

    headers = [
        "Folder Name",
        "Title",
        "Author",
        "Chapters (Audio)",
        "Source Format",
        "Total Duration",
        "Total Duration (sec)",
        "Average Chapter Duration",
        "Total MFA Segments",
        "Segments per Chapter",
        "Average Segment Duration",
        "Matched Units",
        "Status",
        "Missing Units Preview",
        "Folder Path",
    ]
    start_row = 7
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_idx, item in enumerate(stats, start=start_row + 1):
        average_chapter_duration = (
            format_hhmmss(item.total_duration_sec / item.chapter_count)
            if item.chapter_count and item.total_duration_sec > 0
            else "N/V"
        )
        average_segment_duration = (
            format_hhmmss(item.total_duration_sec / item.total_segments)
            if item.total_segments and item.total_duration_sec > 0
            else "N/V"
        )

        values = [
            item.folder_name,
            item.title,
            item.author,
            item.chapter_count,
            item.audio_format,
            format_hhmmss(item.total_duration_sec),
            round(item.total_duration_sec, 3),
            average_chapter_duration,
            item.total_segments,
            round(item.segments_per_chapter, 2) if item.segments_per_chapter is not None else "N/V",
            average_segment_duration,
            item.matched_units,
            item.status,
            item.missing_units_preview,
            str(item.folder_path),
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(row=row_idx, column=col, value=value)

    total_row = start_row + 1 + len(stats)
    ws.cell(row=total_row, column=1, value="TOTAL")
    ws.cell(row=total_row, column=1).font = Font(bold=True)

    # Correct TOTAL row. Text duration columns are derived from the numeric
    # duration-in-seconds column to avoid accidentally summing HH:MM:SS strings.
    ws.cell(row=total_row, column=4, value=f"=SUM(D{start_row + 1}:D{total_row - 1})")
    ws.cell(row=total_row, column=5, value="")
    ws.cell(row=total_row, column=6, value=f'=TEXT(SUM(G{start_row + 1}:G{total_row - 1})/86400,"[h]:mm:ss")')
    ws.cell(row=total_row, column=7, value=f"=SUM(G{start_row + 1}:G{total_row - 1})")
    ws.cell(row=total_row, column=8, value=f'=IF(D{total_row}>0,TEXT(G{total_row}/D{total_row}/86400,"[h]:mm:ss"),"N/V")')
    ws.cell(row=total_row, column=9, value=f"=SUM(I{start_row + 1}:I{total_row - 1})")
    ws.cell(row=total_row, column=10, value=f'=IF(D{total_row}>0,I{total_row}/D{total_row},"N/V")')
    ws.cell(row=total_row, column=11, value=f'=IF(I{total_row}>0,TEXT(G{total_row}/I{total_row}/86400,"[h]:mm:ss"),"N/V")')
    ws.cell(row=total_row, column=12, value=f"=SUM(L{start_row + 1}:L{total_row - 1})")

    # Styling
    thin_gray = Side(style="thin", color="D9E2F3")
    border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)
    for row in ws.iter_rows(min_row=start_row, max_row=total_row, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in ws[total_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E2F0D9")

    # Number formats
    for row in range(start_row + 1, total_row + 1):
        ws.cell(row=row, column=4).number_format = "0"
        ws.cell(row=row, column=7).number_format = "0.000"
        ws.cell(row=row, column=9).number_format = "0"
        ws.cell(row=row, column=10).number_format = "0.00"
        ws.cell(row=row, column=12).number_format = "0"

    # Excel-safe filter range.
    # We intentionally do not create an Excel ListObject/Table here. Some Excel
    # versions repair workbooks when openpyxl tables contain a manually styled
    # total row or formulas. A normal worksheet AutoFilter is sufficient and
    # avoids the /xl/tables/table1.xml repair warning.
    end_col = get_column_letter(len(headers))
    filter_ref = f"A{start_row}:{end_col}{total_row}"

    ws.freeze_panes = f"A{start_row + 1}"
    ws.auto_filter.ref = filter_ref
    autosize_columns(ws)
    ws.column_dimensions["N"].width = 38
    ws.column_dimensions["O"].width = 48

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def validate_dataset_suffix(dataset_suffix: str) -> str:
    """Validate and normalize the selected corpus subset via _00_config.py."""
    try:
        return config.validate_dataset_suffix(dataset_suffix)
    except ValueError as exc:
        raise ValueError(
            f"Invalid DATASET_SUFFIX: {dataset_suffix!r}. "
            "Use 'training', 'validation', or 'test'."
        ) from exc


def run(duration_options: Optional[DurationOptions] = None) -> Path:
    """Run overview generation for DATASET_SUFFIX and return output path."""
    dataset_suffix = validate_dataset_suffix(DATASET_SUFFIX)
    if duration_options is None:
        duration_options = DurationOptions()

    collections_folder = config.dataset_directory(
        config.DATA_COLLECTIONS_DIR / CORPUS_SOURCE_FOLDER,
        dataset_suffix,
    )
    mfa_input_folder = config.dataset_directory(
        config.MFA_INPUT_DIR,
        dataset_suffix,
    )
    output_folder = config.dataset_directory(
        config.CORPUS_OVERVIEW_OUTPUT_DIR,
        dataset_suffix,
    )
    output_path = output_folder / f"corpus_overview_{dataset_suffix}.xlsx"
    cache_path = output_folder / f"audio_duration_cache_{dataset_suffix}.json"

    if not collections_folder.exists() or not collections_folder.is_dir():
        raise SystemExit(
            "Could not find corpus subset folder: "
            f"{collections_folder}"
        )

    if not mfa_input_folder.exists() or not mfa_input_folder.is_dir():
        print(f"WARNING: MFA input folder does not exist yet: {mfa_input_folder}")
        print("Segment counts will be zero / not segmented yet.")

    duration_cache = load_duration_cache(cache_path) if duration_options.use_cache else {}

    all_candidate_folders = sorted(
        [p for p in collections_folder.iterdir() if p.is_dir()],
        key=lambda p: p.name.lower(),
    )
    audiobook_folders = []
    ignored_folders = []
    for folder in all_candidate_folders:
        raw_audio_files = iter_files_by_extensions(folder, AUDIO_EXTENSIONS)
        if raw_audio_files:
            audiobook_folders.append(folder)
        else:
            ignored_folders.append(folder.name)

    if not audiobook_folders:
        raise SystemExit(
            f"No audiobook folders with MP3/WAV files found in: {collections_folder}"
        )

    segment_index = collect_segment_index(mfa_input_folder)
    stats: List[AudiobookStats] = []
    print(f"\nDataset subset: {dataset_suffix}", flush=True)
    print(f"Scanning {len(audiobook_folders)} audiobook folders...", flush=True)
    if not duration_options.calculate_durations:
        print("Duration calculation is skipped for this run.", flush=True)
    elif duration_options.use_cache:
        print(f"Using duration cache: {cache_path}", flush=True)

    for idx, folder in enumerate(audiobook_folders, start=1):
        print(f"  [{idx}/{len(audiobook_folders)}] {folder.name}", flush=True)
        stats.append(
            analyze_audiobook_folder(
                folder,
                segment_index,
                duration_cache,
                duration_options,
            )
        )

    if duration_options.use_cache and duration_options.calculate_durations:
        save_duration_cache(cache_path, duration_cache)

    create_workbook(
        stats,
        collections_folder,
        mfa_input_folder,
        output_path,
        duration_options,
        cache_path,
    )

    print("\nAudiobook corpus overview created.")
    print(f"Dataset subset:     {dataset_suffix}")
    print(f"Audiobooks:         {len(stats)}")
    print(f"Total chapters:     {sum(s.chapter_count for s in stats)}")
    print(f"Total duration:     {format_hhmmss(sum(s.total_duration_sec for s in stats))}")
    print(f"Total MFA segments: {sum(s.total_segments for s in stats)}")
    if ignored_folders:
        print("Ignored folders without supported audio files:")
        for name in ignored_folders:
            print(f"  - {name}")
    print(f"Output:             {output_path}")
    return output_path

def main() -> None:
    parser = argparse.ArgumentParser(description="Create XLSX overview of audiobook corpus metadata and MFA segments.")
    parser.add_argument(
        "--skip-durations",
        action="store_true",
        help="Create a quick overview without calculating audio durations.",
    )
    parser.add_argument(
        "--no-cache",
        action="store_true",
        help="Disable the duration cache for this run.",
    )
    parser.add_argument(
        "--refresh-cache",
        action="store_true",
        help="Recalculate all durations and overwrite cached values.",
    )
    args = parser.parse_args()
    duration_options = DurationOptions(
        calculate_durations=not args.skip_durations,
        use_cache=not args.no_cache,
        refresh_cache=args.refresh_cache,
    )
    run(duration_options)


if __name__ == "__main__":
    main()
